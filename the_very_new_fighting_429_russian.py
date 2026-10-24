#!/usr/bin/env python3
"""
Мониторинг вакансий LinkedIn для русского языка (локализация, копирайтинг, MTPE, LQA и т.п.).
Отдельная версия скрипта: свои настройки, история и таблица.
Находит новые вакансии по поисковым запросам из config_russian.json,
отсеивает уже виденные и нерелевантные, присылает новые в Telegram.
Все найденные вакансии записываются в vacancies_russian.csv рядом со скриптом.

Запуск:  python linkedin_jobs_russian.py
Зависимости:  pip install requests beautifulsoup4
"""

import csv
import json
import os
import random
import re
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import quote_plus, urlencode

import requests
from bs4 import BeautifulSoup

try:
    from docx import Document          # pip install python-docx
except ImportError:
    Document = None
try:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    pass

BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = BASE_DIR / ".env"                  # TG_TOKEN и TG_CHAT_ID
CONFIG_PATH = BASE_DIR / "config_russian.json"
SEEN_PATH = BASE_DIR / "seen_jobs_russian.json"
PENDING_PATH = BASE_DIR / "pending_jobs_russian.json"   # очередь непроверенных описаний
LOG_PATH = BASE_DIR / "vacancies_russian.csv"           # таблица всех найденных вакансий

SEARCH_URL = "https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search"
DETAIL_URL = "https://www.linkedin.com/jobs-guest/jobs/api/jobPosting/{}"
PAGE_SIZE = 10          # гостевой поиск отдаёт около 10 вакансий за запрос
SEEN_LIMIT = 5000       # сколько ID вакансий помнить
DOCX_PATH = BASE_DIR / "descriptions_russian.docx"   # описания вакансий текущего прогона
MAX_ATTEMPTS = 3        # сколько раз пытаться открыть описание, потом сдаться

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}


class RateLimited(Exception):
    """LinkedIn ответил 429 — пора остановиться."""


# ---------- вспомогательное ----------

def load_env(path=ENV_PATH):
    """
    Читает строки вида KEY=VALUE из файла .env в os.environ.
    Если переменная уже задана в системе, системное значение остаётся главным.
    """
    try:
        lines = path.read_text(encoding="utf-8-sig").splitlines()
    except FileNotFoundError:
        return False
    for line in lines:
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        if key.startswith("export "):
            key = key[len("export "):].strip()
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)
    return True


def telegram_credentials(cfg):
    """Токен и chat_id: из .env (TG_TOKEN, TG_CHAT_ID), запасной вариант — старые поля config."""
    token = os.environ.get("TG_TOKEN") or cfg.get("telegram_bot_token", "")
    chat = os.environ.get("TG_CHAT_ID") or cfg.get("telegram_chat_id", "")
    return token.strip(), str(chat).strip()


def load_json(path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def save_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def pause(cfg):
    lo, hi = cfg.get("delay_between_requests_sec", [4, 9])
    time.sleep(random.uniform(lo, hi))


def pause_desc(cfg):
    lo, hi = cfg.get("description_delay_sec", [2, 4])
    time.sleep(random.uniform(lo, hi))


def find_word(words, text):
    """
    Ищет слова из списка в тексте. Слово должно стоять целиком
    (допускается окончание -s): 'crypto' найдёт 'crypto' и 'cryptos',
    но не 'cryptography'; 'defi' не сработает на 'definitely'.
    Возвращает первое найденное слово или None.
    """
    text = text.lower()
    for w in words:
        pattern = r"(?<![a-z0-9])" + re.escape(w.lower()) + r"s?(?![a-z0-9])"
        if re.search(pattern, text):
            return w
    return None


# ---------- поиск и разбор ----------

def build_url(keywords, cfg, start):
    params = {
        "keywords": keywords,
        "location": cfg.get("location", "Worldwide"),
        "sortBy": "DD",                                         # сначала новые
        "f_TPR": f"r{int(cfg.get('posted_within_hours', 24)) * 3600}",
        "start": start,
    }
    if cfg.get("remote_only"):
        params["f_WT"] = "2"                                    # 2 = удалёнка
    return f"{SEARCH_URL}?{urlencode(params)}"


def parse_jobs(html):
    soup = BeautifulSoup(html, "html.parser")
    jobs = []
    for card in soup.select("div.base-card, div.job-search-card"):
        urn = card.get("data-entity-urn", "")
        job_id = urn.rsplit(":", 1)[-1] if urn else ""

        def text(sel):
            el = card.select_one(sel)
            return el.get_text(" ", strip=True) if el else ""

        link_el = card.select_one("a.base-card__full-link, a[href*='/jobs/view/']")
        link = link_el["href"].split("?")[0] if link_el and link_el.get("href") else ""
        if not job_id and link:
            job_id = link.rstrip("/").rsplit("-", 1)[-1]
        if not job_id:
            continue

        time_el = card.select_one("time")
        jobs.append({
            "id": job_id,
            "title": text("h3.base-search-card__title, h3"),
            "company": text("h4.base-search-card__subtitle, h4"),
            "location": text("span.job-search-card__location"),
            "posted": time_el.get_text(strip=True) if time_el else "",
            "posted_date": time_el.get("datetime", "") if time_el else "",
            "link": link or f"https://www.linkedin.com/jobs/view/{job_id}/",
        })
    return jobs


def linkedin_get(session, url, cfg):
    """
    Запрос к LinkedIn. При 429 ждёт и повторяет (паузы из rate_limit_waits_sec).
    Если и после ожиданий 429 — выбрасывает RateLimited. None — сетевая ошибка.
    """
    waits = cfg.get("rate_limit_waits_sec", [60, 180])
    for attempt in range(len(waits) + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=20)
        except requests.RequestException as e:
            print(f"  ! сеть: {e}", flush=True)
            return None
        if resp.status_code != 429:
            return resp
        if attempt == len(waits):
            break
        wait = int(waits[attempt])
        retry_after = resp.headers.get("Retry-After", "")
        if retry_after.isdigit():
            wait = max(wait, min(int(retry_after), 900))
        print(f"  ! LinkedIn просит притормозить (429) — жду {wait} сек и пробую снова...", flush=True)
        time.sleep(wait)
    raise RateLimited()


def fetch_search(session, keywords, cfg):
    found = []
    for page in range(int(cfg.get("pages_per_search", 5))):
        url = build_url(keywords, cfg, page * PAGE_SIZE)
        resp = linkedin_get(session, url, cfg)          # при стойком 429 — RateLimited
        if resp is None:
            break
        if resp.status_code != 200:
            print(f"  ! ответ {resp.status_code}, пропускаю")
            break
        jobs = parse_jobs(resp.text)
        found.extend(jobs)
        if len(jobs) < PAGE_SIZE:
            break                                               # дальше страниц нет
        pause(cfg)
    return found


def parse_description(html):
    """Текст блока «About the job» с абзацами и пунктами списков."""
    soup = BeautifulSoup(html, "html.parser")
    el = soup.select_one(
        "div.show-more-less-html__markup, div.description__text, section.description"
    )
    if not el:
        return None
    for li in el.find_all("li"):
        li.insert(0, "• ")
    for br in el.find_all("br"):
        br.replace_with("\n")
    for tag in el.find_all(["p", "li", "div", "ul", "ol", "h1", "h2", "h3", "h4"]):
        tag.insert_before("\n")
        tag.insert_after("\n")
    lines = [re.sub(r"\s+", " ", line).strip() for line in el.get_text("").split("\n")]
    text = "\n".join(line for line in lines if line)
    return text or None


def fetch_description(session, job_id, cfg):
    """Текст описания вакансии или None, если получить не удалось."""
    resp = linkedin_get(session, DETAIL_URL.format(job_id), cfg)   # при стойком 429 — RateLimited
    if resp is None:
        return None
    if resp.status_code != 200:
        print(f"  ! описание {job_id}: ответ {resp.status_code}")
        return None
    return parse_description(resp.text)


# ---------- Telegram ----------

def send_telegram(cfg, text):
    token, chat = telegram_credentials(cfg)
    if not token or not chat:
        print(text + "\n")
        return
    try:
        r = requests.post(
            f"https://api.telegram.org/bot{token}/sendMessage",
            data={"chat_id": chat, "text": text, "parse_mode": "HTML",
                  "disable_web_page_preview": "true"},
            timeout=20,
        )
        if r.status_code != 200:
            print(f"  ! Telegram: {r.status_code} {r.text[:200]}")
    except requests.RequestException as e:
        print(f"  ! Telegram: {e}")
    time.sleep(1)


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def people_links(company, groups):
    """
    Для компании и групп должностей строит ссылки на поиск людей:
    [(название группы, ссылка LinkedIn, ссылка Google X-ray), ...]
    Никаких запросов не делает — только собирает ссылки.
    """
    company = company.replace('"', "").strip()
    if not company or not groups:
        return []
    links = []
    for group, roles in groups.items():
        roles_q = " OR ".join(f'"{r}"' if " " in r else r for r in roles)
        li = ("https://www.linkedin.com/search/results/people/?keywords="
              + quote_plus(f'"{company}" AND ({roles_q})'))
        google = ("https://www.google.com/search?q="
                  + quote_plus(f'site:linkedin.com/in "{company}" ({roles_q})'))
        links.append((group, li, google))
    return links


def format_job(job, query, tag="", groups=None):
    lines = [f"{esc(tag)} <b>{esc(job['title'])}</b>" if tag else f"<b>{esc(job['title'])}</b>", esc(job["company"])]
    meta = " · ".join(x for x in (job["location"], posted_date_str(job) or job["posted"]) if x)
    if meta:
        lines.append(meta)
    lines.append(f"🔎 {esc(query)}")
    if job.get("matched_in_description"):
        lines.append(f"📄 в описании: {esc(job['matched_in_description'])}")
    lines.append(job["link"])
    people = people_links(job.get("company", ""), groups)
    if people:
        lines.append("")
        lines.append("<b>Кому написать:</b>")
        for group, li, google in people:
            lines.append(f'👥 {esc(group)}: <a href="{esc(li)}">LinkedIn</a> · <a href="{esc(google)}">Google</a>')
    return "\n".join(lines)


# ---------- дата публикации и файл с описаниями ----------

def posted_date_str(job):
    """Точная дата публикации в виде ДД.ММ.ГГГГ (или пусто)."""
    d = job.get("posted_date", "")
    if d:
        try:
            return datetime.strptime(d[:10], "%Y-%m-%d").strftime("%d.%m.%Y")
        except ValueError:
            pass
    m = re.search(r"(\d+)\s*(minute|hour|day|week|month)", job.get("posted", "").lower())
    if m:
        n, unit = int(m.group(1)), m.group(2)
        delta = {"minute": timedelta(minutes=n), "hour": timedelta(hours=n),
                 "day": timedelta(days=n), "week": timedelta(weeks=n),
                 "month": timedelta(days=30 * n)}[unit]
        return (datetime.now() - delta).strftime("%d.%m.%Y")
    return ""


def add_hyperlink(paragraph, url, text):
    """Кликабельная ссылка в docx (в python-docx нет готовой функции)."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def write_docx(entries, started, groups=None):
    """Сохраняет описания вакансий этого прогона, заменяя прошлый файл."""
    if Document is None:
        print("  ! Не установлен python-docx (pip install python-docx) — файл с описаниями не создан", flush=True)
        return
    doc = Document()
    doc.add_heading(f"Вакансии LinkedIn — прогон {started:%d.%m.%Y %H:%M}", level=1)
    if not entries:
        doc.add_paragraph("В этом прогоне новых подходящих вакансий нет.")
    for n, (job, desc) in enumerate(entries, 1):
        doc.add_heading(f"Вакансия № {n}", level=2)
        for label, value in (("Вакансия", job.get("title", "")),
                             ("Компания", job.get("company", "")),
                             ("Дата публикации", posted_date_str(job) or "не указана"),
                             ("Ссылка", job.get("link", ""))):
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(value)
        people = people_links(job.get("company", ""), groups)
        if people:
            doc.add_paragraph().add_run("Кому написать:").bold = True
            for group, li, google in people:
                p = doc.add_paragraph(f"{group}: ")
                add_hyperlink(p, li, "LinkedIn")
                p.add_run(" · ")
                add_hyperlink(p, google, "Google")
        doc.add_paragraph().add_run("About the job:").bold = True
        if desc:
            for line in desc.split("\n"):
                doc.add_paragraph(line)
        else:
            doc.add_paragraph().add_run("Описание получить не удалось.").italic = True

    tmp = DOCX_PATH.with_name(DOCX_PATH.stem + "_tmp.docx")
    doc.save(tmp)
    try:
        os.replace(tmp, DOCX_PATH)
        print(f"Описания ({len(entries)}) сохранены в {DOCX_PATH.name}", flush=True)
    except PermissionError:
        alt = DOCX_PATH.with_name(f"{DOCX_PATH.stem}_{started:%Y-%m-%d_%H-%M}.docx")
        os.replace(tmp, alt)
        print(f"  ! {DOCX_PATH.name} открыт в Word — описания сохранены в {alt.name}", flush=True)


# ---------- таблица вакансий ----------

LOG_FIELDS = ["Дата проверки", "Статус", "Причина", "Название", "Компания",
              "Локация", "Опубликовано", "Запрос", "Ссылка"]


def log_job(job, query, status, reason):
    """Дописывает строку в vacancies_russian.csv (открывается в Excel)."""
    new_file = not LOG_PATH.exists()
    try:
        with LOG_PATH.open("a", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            if new_file:
                w.writerow(LOG_FIELDS)
            w.writerow([datetime.now().strftime("%Y-%m-%d %H:%M"), status, reason,
                        job.get("title", ""), job.get("company", ""), job.get("location", ""),
                        job.get("posted", ""), query, job.get("link", "")])
    except PermissionError:
        print("  ! vacancies_russian.csv открыт в другой программе (Excel?) — строка не записана", flush=True)


# ---------- main ----------

def main():
    cfg = load_json(CONFIG_PATH, None)
    if cfg is None:
        sys.exit(f"Не найден или повреждён {CONFIG_PATH}")

    if not load_env():
        print(f"  ! Не найден файл {ENV_PATH.name} рядом со скриптом", flush=True)
    token, chat = telegram_credentials(cfg)
    if not token or not chat:
        print("  ! TG_TOKEN или TG_CHAT_ID не заданы в .env — вакансии будут выводиться в консоль, "
              "а не в Telegram", flush=True)

    exclude = cfg.get("exclude_words", [])
    require = cfg.get("require_any_words", [])
    roles = cfg.get("require_role_words", [])     # профессия — проверяется только в названии
    check_desc = cfg.get("check_description", True)
    max_checks = int(cfg.get("max_description_checks", 60))

    seen = load_json(SEEN_PATH, [])
    seen_set = set(seen)
    first_run = not seen
    session = requests.Session()

    def mark_seen(job_id):
        if job_id not in seen_set:
            seen_set.add(job_id)
            seen.append(job_id)

    # Очередь на проверку описаний: сначала отложенные с прошлых запусков
    queue = []            # элементы: {"job": {...}, "query": "...", "attempts": N}
    queued_ids = set()
    for item in load_json(PENDING_PATH, []):
        jid = item.get("job", {}).get("id")
        if jid and jid not in seen_set and jid not in queued_ids:
            queue.append(item)
            queued_ids.add(jid)
    from_pending = len(queue)

    results = []          # (job, query) — пойдут в Telegram
    this_run = set()
    dropped = 0

    # 1. Поиск по всем запросам
    search_blocked = False
    for query in cfg.get("searches", []):
        print(f"Ищу: {query}", flush=True)
        try:
            found = fetch_search(session, query, cfg)
        except RateLimited:
            search_blocked = True
            print("  ! LinkedIn ограничил запросы — прекращаю поиск, обработаю то, что успел найти", flush=True)
            break
        for job in found:
            jid = job["id"]
            if jid in seen_set or jid in this_run or jid in queued_ids:
                continue
            this_run.add(jid)
            head = f"{job['title']} {job['company']}"

            stop = find_word(exclude, head)
            if stop:                                      # стоп-слово в названии/компании
                mark_seen(jid)
                dropped += 1
                log_job(job, query, "отсеяна", f"стоп-слово: {stop}")
                continue
            if roles and not find_word(roles, job["title"]):  # не контентная профессия
                mark_seen(jid)
                dropped += 1
                log_job(job, query, "отсеяна", "не та профессия (нет слов из require_role_words в названии)")
                continue
            word = find_word(require, head) if require else "фильтр не задан"
            if word:                                      # нужное слово уже в названии/компании
                results.append((job, query, word))        # отправим после чтения описания
            elif check_desc:                              # проверим текст вакансии
                queue.append({"job": job, "query": query, "attempts": 0})
                queued_ids.add(jid)
            else:
                mark_seen(jid)
                dropped += 1
                log_job(job, query, "отсеяна", "нет нужных слов в названии")
        pause(cfg)

    started = datetime.now()
    sent = 0
    retry = []            # не открылись в этот раз — попробуем снова
    pos = 0               # сколько элементов очереди уже обработано
    checks = 0            # сколько описаний открыто за этот запуск
    blocked = search_blocked   # LinkedIn ответил 429 — больше к нему не обращаемся
    docx_entries = []     # (job, описание) — для файла с описаниями

    def save():
        save_json(SEEN_PATH, seen[-SEEN_LIMIT:])
        save_json(PENDING_PATH, retry + queue[pos:])

    def notify(job, query):
        nonlocal sent, first_run
        if first_run:
            send_telegram(cfg, "Мониторинг LinkedIn (русский язык) запущен. Присылаю найденные вакансии:")
            first_run = False
        send_telegram(cfg, format_job(job, query, cfg.get("message_tag", ""), cfg.get("people_search_groups")))
        sent += 1

    def deliver(job, query, desc, reason):
        mark_seen(job["id"])
        log_job(job, query, "отправлена", reason)
        notify(job, query)
        docx_entries.append((job, desc))

    def open_description(job):
        """Открывает описание с учётом лимита и паузы. None — не получилось."""
        nonlocal checks, blocked
        if blocked or checks >= max_checks:
            return None
        if checks:
            pause_desc(cfg)
        checks += 1
        try:
            return fetch_description(session, job["id"], cfg)
        except RateLimited:
            blocked = True
            print("  ! LinkedIn ограничил запросы — остальные описания проверю в следующий раз", flush=True)
            return None

    print(f"Прошли по названию: {len(results)} | на проверку описаний: {len(queue)}"
          + (f" (из них отложенных с прошлого раза: {from_pending})" if from_pending else ""),
          flush=True)
    try:
        # 2. Прошедшие по названию: читаем описание (для файла) и сразу отправляем
        if results:
            print("Читаю описания вакансий, прошедших по названию...", flush=True)
        for n, (job, query, word) in enumerate(results, 1):
            desc = open_description(job)
            note = "" if desc else " (описание не получено)"
            print(f"  [{n}/{len(results)}] {job['title'][:50]} — {job['company'][:30]} → ✅ {word}{note}", flush=True)
            deliver(job, query, desc, f"в названии/компании: {word}")
            save()

        # 3. Остальные: проверяем описание, подходящие отправляем
        total = min(len(queue), max(0, max_checks - checks)) if not blocked else 0
        if queue:
            print(f"Проверяю описания ({total} из {len(queue)}), это займёт несколько минут...", flush=True)
        n = 0
        while pos < len(queue) and checks < max_checks and not blocked:
            item = queue[pos]
            job, query = item["job"], item["query"]
            label = f"  [{n + 1}/{total}] {job['title'][:50]} — {job['company'][:30]}"
            desc = open_description(job)
            if blocked:
                break
            pos += 1
            n += 1
            if desc is None:
                item["attempts"] = item.get("attempts", 0) + 1
                if item["attempts"] >= MAX_ATTEMPTS:
                    mark_seen(job["id"])
                    dropped += 1
                    log_job(job, query, "отсеяна", f"описание не открылось {MAX_ATTEMPTS} раза")
                    print(f"{label} → не открылось {MAX_ATTEMPTS} раза, пропускаю", flush=True)
                else:
                    retry.append(item)
                    print(f"{label} → не открылось, попробую в следующий раз", flush=True)
            else:
                word = find_word(require, desc)
                if word:
                    job["matched_in_description"] = word
                    print(f"{label} → ✅ {word}", flush=True)
                    deliver(job, query, desc, f"в описании: {word}")
                else:
                    mark_seen(job["id"])
                    dropped += 1
                    log_job(job, query, "отсеяна", "нет нужных слов в описании")
                    print(f"{label} → мимо", flush=True)
            save()                                        # прогресс не теряется при прерывании
    except KeyboardInterrupt:
        print("\nОстановлено вручную. Прогресс сохранён.", flush=True)
    finally:
        save()
        write_docx(docx_entries, started, cfg.get("people_search_groups"))

    postponed = len(retry) + len(queue) - pos
    if blocked:
        send_telegram(cfg, "⚠️ LinkedIn временно ограничил запросы (429), прогон остановлен досрочно.\n"
                           f"Отправлено вакансий: {sent}. Непроверенное скрипт доделает в следующий запуск.\n"
                           "Если это повторяется — запускайте реже и разнесите запуски двух скриптов по времени.")
    print(f"Готово. Отправлено в Telegram: {sent} | отсеяно: {dropped}"
          + (f" | в очереди на следующий запуск: {postponed}" if postponed else ""), flush=True)


if __name__ == "__main__":
    main()